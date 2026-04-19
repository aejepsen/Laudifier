# backend/pipeline/run_pipeline.py
"""
Pipeline de indexação de laudos no Qdrant.
Funções de chunking e extração de texto reutilizadas pelo agente de busca.
"""

import re
import os
import uuid
import asyncio
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

_CHUNK_SIZE    = 1500   # chars por chunk
_CHUNK_OVERLAP = 150


# ─── Extração de texto ────────────────────────────────────────────────────────

def _extrair_texto(filepath: str, filename: str) -> str:
    """Extrai texto de um arquivo (.txt, .pdf, .docx)."""
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        return Path(filepath).read_text(encoding="utf-8", errors="ignore")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        return "\n".join(p.extract_text() or "" for p in reader.pages)

    if ext in (".docx", ".doc"):
        from docx import Document
        doc = Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs)

    return Path(filepath).read_text(encoding="utf-8", errors="ignore")


# ─── Anonimização ─────────────────────────────────────────────────────────────

def _anonimizar_texto(texto: str) -> str:
    """
    Remove dados identificadores do paciente antes de indexar no Qdrant.
    Cobre: nomes, CPF, datas, CRM, médico, placeholders existentes.
    """
    padroes = [
        # Nome/Paciente em linha de campo
        (r'\b(?:Paciente|Nome)\s*:\s*[^\n]+',                     'Paciente: [PACIENTE]'),
        # CPF (com ou sem pontuação)
        (r'\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b',                    '[CPF]'),
        # Datas DD/MM/AAAA e variações
        (r'\b\d{1,2}/\d{1,2}/\d{2,4}\b',                         '[DATA]'),
        # CRM com código
        (r'\bCRM\s*[:\-]?\s*[\w\-/]+',                           'CRM: [CRM]'),
        # Dr./Dra. seguido de nome próprio
        (r'\bDr[aA]?\.?\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][a-záéíóúâêôãõç]+'
         r'(?:\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][a-záéíóúâêôãõç]+)*',          'Dr. [MÉDICO]'),
        # Placeholders já existentes no template
        (r'\[NOME DO PACIENTE\]',                                  '[PACIENTE]'),
        (r'\[DATA DO EXAME\]',                                     '[DATA]'),
        (r'\[DATA DE NASCIMENTO\]',                                '[DATA]'),
        (r'\[CRM DO MÉDICO\]',                                     '[CRM]'),
        (r'\[ASSINATURA(?: DO MÉDICO)?\]',                         '[MÉDICO]'),
    ]
    for pattern, replacement in padroes:
        texto = re.sub(pattern, replacement, texto, flags=re.IGNORECASE)
    return texto


# ─── Chunking ─────────────────────────────────────────────────────────────────

def _chunk_laudo(
    texto: str,
    filename: str,
    especialidade: str,
    tipo_laudo: str,
    user_id: str,
) -> list[dict]:
    """
    Divide o texto do laudo em chunks com metadados para indexação no Qdrant.
    Laudos curtos retornam um único chunk; laudos longos são divididos por seção
    ou por tamanho quando não há seções explícitas.
    """
    if not texto.strip():
        return []

    section_pattern = re.compile(r'^([A-ZÁÉÍÓÚÂÊÔÃÕÇ\s/:]+):\s*$', re.MULTILINE)
    sections = section_pattern.split(texto.strip())

    chunks: list[str] = []

    if len(sections) > 1:
        i = 1
        while i < len(sections) - 1:
            title   = sections[i].strip()
            content = sections[i + 1].strip()
            if content:
                chunks.append(f"{title}:\n{content}")
            i += 2
    else:
        start = 0
        while start < len(texto):
            end = min(start + _CHUNK_SIZE, len(texto))
            chunks.append(texto[start:end])
            if end == len(texto):
                break
            start = end - _CHUNK_OVERLAP

    if not chunks:
        chunks = [texto.strip()]

    return [
        {
            "texto":         chunk,
            "filename":      filename,
            "especialidade": especialidade,
            "tipo_laudo":    tipo_laudo,
            "user_id":       user_id,
            "aprovado":      True,
        }
        for chunk in chunks
        if chunk.strip()
    ]


# ─── Indexação no Qdrant ──────────────────────────────────────────────────────

async def _indexar_chunks_repositorio(
    chunks: list[dict],
    filename: str,
    especialidade: str,
    tipo_laudo: str,
    user_id: str,
) -> None:
    """Converte chunks em vetores e upserta no Qdrant como fonte 'repositorio'."""
    from qdrant_client.models import PointStruct
    from ..agents.search_agent import _get_model, _get_qdrant_client

    COLLECTION = os.getenv("QDRANT_COLLECTION", "laudos_medicos")

    model  = _get_model()
    qdrant = _get_qdrant_client()

    points = []
    for i, chunk in enumerate(chunks):
        texto = chunk["texto"]
        vec = await asyncio.to_thread(
            model.encode,
            f"passage: {texto}",
            normalize_embeddings=True,
        )
        point_id = str(uuid.uuid5(
            uuid.NAMESPACE_DNS,
            f"repositorio:{user_id}:{filename}:{i}",
        ))
        points.append(PointStruct(
            id=point_id,
            vector=vec.tolist(),
            payload={
                "content":       texto,
                "source_name":   filename,
                "especialidade": especialidade.lower(),
                "tipo_laudo":    tipo_laudo,
                "source":        "repositorio",
                "user_id":       user_id,
                "chunk_index":   i,
            },
        ))

    await qdrant.upsert(collection_name=COLLECTION, points=points)


# ─── Ponto de entrada público ─────────────────────────────────────────────────

async def ingerir_laudo(
    url: str,
    filename: str,
    especialidade: str,
    tipo_laudo: str,
    user_id: str,
) -> None:
    """
    Pipeline completo de ingestão de laudo no repositório:
      1. Lê o arquivo (caminho local ou URL HTTP)
      2. Extrai texto (PDF / DOCX / TXT)
      3. Anonimiza dados pessoais do paciente
      4. Divide em chunks por seção
      5. Indexa no Qdrant como fonte 'repositorio'

    Dados do paciente (nome, CPF, datas, CRM) são removidos antes de qualquer
    persistência no banco vetorial — conformidade com LGPD art. 18.
    """
    import tempfile

    filepath = None
    tmp_created = False

    try:
        if url.startswith("http"):
            import httpx
            async with httpx.AsyncClient(timeout=30) as client:
                resp = await client.get(url)
                resp.raise_for_status()
                content = resp.content
            suffix = Path(filename).suffix or ".bin"
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(content)
                filepath = tmp.name
                tmp_created = True
        else:
            filepath = url  # caminho local (USE_LOCAL_STORAGE=true)

        texto = _extrair_texto(filepath, filename)
        if not texto.strip():
            logger.warning("[Ingestão] Arquivo sem texto extraível: %s", filename)
            return

        texto_anonimizado = _anonimizar_texto(texto)
        chunks = _chunk_laudo(texto_anonimizado, filename, especialidade, tipo_laudo, user_id)

        if not chunks:
            logger.warning("[Ingestão] Nenhum chunk gerado: %s", filename)
            return

        await _indexar_chunks_repositorio(chunks, filename, especialidade, tipo_laudo, user_id)
        logger.info("[Ingestão] %s indexado — %d chunks, dados pessoais anonimizados", filename, len(chunks))

    except Exception:
        logger.error("[Ingestão] Falhou para %s", filename, exc_info=True)
        raise
    finally:
        if tmp_created and filepath and Path(filepath).exists():
            Path(filepath).unlink()
