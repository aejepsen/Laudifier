"""Cobertura de ExportService — PDF/DOCX/TXT."""
from __future__ import annotations

import os

import pytest

from backend.services.export_service import ExportService


LAUDO_BASE = {
    "laudo": "ACHADOS:\nCampos pulmonares livres.\n\nIMPRESSÃO:\nNormal.",
    "laudo_editado": None,
    "especialidade": "radiologia",
    "created_at": "2026-04-24T10:00:00",
}


# ── dispatch ─────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_exportar_pdf_dispatch():
    svc = ExportService()
    path = await svc.exportar(LAUDO_BASE, "pdf")
    assert path.endswith(".pdf")
    assert os.path.getsize(path) > 0


@pytest.mark.asyncio
async def test_exportar_docx_dispatch():
    svc = ExportService()
    path = await svc.exportar(LAUDO_BASE, "docx")
    assert path.endswith(".docx")
    assert os.path.getsize(path) > 0


@pytest.mark.asyncio
async def test_exportar_txt_dispatch():
    svc = ExportService()
    path = await svc.exportar(LAUDO_BASE, "txt")
    assert path.endswith(".txt")
    with open(path, encoding="utf-8") as f:
        content = f.read()
    assert "RADIOLOGIA" in content
    assert "ACHADOS" in content


@pytest.mark.asyncio
async def test_exportar_formato_desconhecido_cai_em_txt():
    svc = ExportService()
    path = await svc.exportar(LAUDO_BASE, "qualquercoisa")
    assert path.endswith(".txt")


# ── precedência laudo_editado > laudo ─────────────────────────────────────────

@pytest.mark.asyncio
async def test_laudo_editado_tem_precedencia():
    svc = ExportService()
    laudo = {
        **LAUDO_BASE,
        "laudo_editado": "VERSÃO EDITADA PELO MÉDICO",
    }
    path = await svc.exportar(laudo, "txt")
    with open(path, encoding="utf-8") as f:
        content = f.read()
    assert "VERSÃO EDITADA" in content
    assert "Campos pulmonares" not in content


@pytest.mark.asyncio
async def test_laudo_editado_vazio_usa_laudo_original():
    svc = ExportService()
    laudo = {**LAUDO_BASE, "laudo_editado": ""}
    path = await svc.exportar(laudo, "txt")
    with open(path, encoding="utf-8") as f:
        content = f.read()
    assert "Campos pulmonares" in content


# ── PDF: branches de formatação (uppercase headers, **, linha vazia) ─────────

def test_pdf_renderiza_headers_uppercase_e_corpo():
    svc = ExportService()
    laudo = {
        **LAUDO_BASE,
        "laudo": "ACHADOS\nLinha de corpo com **negrito**.\n\nIMPRESSÃO\nNormal.",
    }
    path = svc._to_pdf(laudo["laudo"], laudo)
    assert os.path.exists(path)
    assert os.path.getsize(path) > 500  # PDF não vazio


def test_pdf_linha_uppercase_longa_nao_vira_heading():
    """Linha uppercase >= 60 chars cai no ramo body_style (não Heading2)."""
    svc = ExportService()
    longa = "A" * 70  # uppercase mas comprida
    laudo = {**LAUDO_BASE, "laudo": f"{longa}\nCorpo."}
    path = svc._to_pdf(laudo["laudo"], laudo)
    assert os.path.exists(path)


def test_pdf_sem_especialidade_nem_data():
    svc = ExportService()
    laudo = {"laudo": "Conteúdo mínimo."}
    path = svc._to_pdf(laudo["laudo"], laudo)
    assert os.path.exists(path)


# ── DOCX: branches simétricos ────────────────────────────────────────────────

def test_docx_renderiza_headers_e_corpo():
    svc = ExportService()
    path = svc._to_docx(LAUDO_BASE["laudo"], LAUDO_BASE)
    assert os.path.exists(path)
    assert os.path.getsize(path) > 0

    from docx import Document
    doc = Document(path)
    paragrafos = [p.text for p in doc.paragraphs]
    assert any("ACHADOS" in p for p in paragrafos)
    assert any("Campos pulmonares" in p for p in paragrafos)


def test_docx_remove_marcadores_negrito():
    svc = ExportService()
    laudo = {**LAUDO_BASE, "laudo": "Linha com **negrito** dentro."}
    path = svc._to_docx(laudo["laudo"], laudo)
    from docx import Document
    doc = Document(path)
    texto_total = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in texto_total
    assert "negrito" in texto_total


# ── TXT: encoding utf-8 e estrutura ──────────────────────────────────────────

def test_txt_encoding_utf8_preserva_acentos():
    svc = ExportService()
    laudo = {**LAUDO_BASE, "laudo": "Avaliação de coração e pulmões."}
    path = svc._to_txt(laudo["laudo"], laudo)
    with open(path, encoding="utf-8") as f:
        content = f.read()
    assert "Avaliação" in content
    assert "coração" in content
